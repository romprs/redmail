"""Сборка инструкции пользователя из docs/УСТАНОВКА_И_НАСТРОЙКА.md в .docx.

Запуск: python docs/build_manual.py [каталог назначения]
По умолчанию .docx кладётся рядом с исходным .md.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

SOURCE = Path(__file__).with_name("УСТАНОВКА_И_НАСТРОЙКА.md")


def _add_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for col_index, cell_text in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = cell_text
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.bold = row_index == 0


def _split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def build(md_path: Path = SOURCE, out_dir: Path | None = None) -> Path:
    text = md_path.read_text(encoding="utf-8")
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "DejaVu Sans"
    style.font.size = Pt(11)

    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            index += 1
            code: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            paragraph = document.add_paragraph()
            run = paragraph.add_run("\n".join(code))
            run.font.name = "DejaVu Sans Mono"
            run.font.size = Pt(9)
            paragraph.paragraph_format.left_indent = Pt(18)
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and set(lines[index + 1].strip()) <= set("|-: "):
            rows = [_split_table_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_split_table_row(lines[index]))
                index += 1
            _add_table(document, rows)
            document.add_paragraph()
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            heading = document.add_heading(stripped.lstrip("# ").strip(), level=min(level, 3))
            if level == 1:
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith(("- ", "* ")):
            document.add_paragraph(_plain(stripped[2:]), style="List Bullet")
        elif stripped:
            document.add_paragraph(_plain(stripped))
        index += 1

    target_dir = out_dir or md_path.parent
    target_dir.mkdir(parents=True, exist_ok=True)
    out_path = target_dir / (md_path.stem + ".docx")
    document.save(out_path)
    return out_path


def _plain(value: str) -> str:
    """Разметка markdown внутри строки в обычный текст."""
    value = re.sub(r"\*\*(.+?)\*\*", r"\1", value)
    value = re.sub(r"`([^`]+)`", r"\1", value)
    return value


if __name__ == "__main__":
    destination = Path(sys.argv[1]) if len(sys.argv) > 1 else None
    print(build(out_dir=destination))
